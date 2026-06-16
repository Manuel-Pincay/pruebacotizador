from openpyxl import load_workbook
from openpyxl import Workbook

from openpyxl.styles import Font
from openpyxl.styles import PatternFill
from openpyxl.styles import Alignment

from openpyxl.utils import get_column_letter

from app.models.client import Client
from app.models.product import Product


# ==========================================
# STYLES
# ==========================================

HEADER_FILL = PatternFill(
    start_color="7C3AED",
    end_color="7C3AED",
    fill_type="solid"
)

HEADER_FONT = Font(
    color="FFFFFF",
    bold=True
)

DESCRIPTION_FILL = PatternFill(
    start_color="E9D5FF",
    end_color="E9D5FF",
    fill_type="solid"
)

CENTER = Alignment(
    horizontal="center",
    vertical="center"
)


# ==========================================
# FORMAT SHEET
# ==========================================

def format_sheet(ws):

    # HEADER STYLE

    for cell in ws[1]:

        cell.fill = HEADER_FILL

        cell.font = HEADER_FONT

        cell.alignment = CENTER

    # DESCRIPTION STYLE

    for cell in ws[2]:

        cell.fill = DESCRIPTION_FILL

        cell.alignment = CENTER

    # AUTO WIDTH

    for column_cells in ws.columns:

        length = max(

            len(str(cell.value or ""))

            for cell in column_cells

        )

        col_letter = get_column_letter(
            column_cells[0].column
        )

        ws.column_dimensions[
            col_letter
        ].width = length + 8

    # FREEZE HEADER

    ws.freeze_panes = "A2"


# ==========================================
# CLIENT TEMPLATE
# ==========================================

def create_clients_template(path):

    wb = Workbook()

    ws = wb.active

    ws.title = "Clientes"

    # HEADERS

    headers = [

        "name",
        "company",
        "ruc_ci",
        "phone",
        "email",
        "address",
        "client_type",
        "observations"

    ]

    ws.append(headers)

    # DESCRIPTIONS

    ws.append([

        "Nombre cliente",
        "Empresa",
        "RUC o CI",
        "Teléfono",
        "Correo",
        "Dirección",
        "Mayorista / Minorista",
        "Notas adicionales"

    ])

    # EXAMPLE

    ws.append([

        "Juan Pérez",
        "Innova Arte",
        "1723456789",
        "0999999999",
        "juan@gmail.com",
        "Quito",
        "Mayorista",
        "Cliente frecuente"

    ])

    format_sheet(ws)

    wb.save(path)


# ==========================================
# PRODUCT TEMPLATE
# ==========================================

def create_products_template(path):

    wb = Workbook()

    ws = wb.active

    ws.title = "Productos"

    # HEADERS

    headers = [

        "code",
        "name",
        "description",
        "category",
        "material",
        "color",
        "size",
        "thickness",
        "price",
        "cost",
        "theme",
        "stock",
        "custom"

    ]

    ws.append(headers)

    # DESCRIPTIONS

    ws.append([

        "Código producto",
        "Nombre producto",
        "Descripción",
        "Categoría",
        "Material",
        "Color",
        "Tamaño",
        "Espesor",
        "Precio venta",
        "Costo",
        "Temática",
        "Stock",
        "TRUE/FALSE"

    ])

    # EXAMPLE

    ws.append([

        "TOP001",
        "Topper Feliz Cumpleaños",
        "Topper MDF dorado",
        "Topper",
        "MDF",
        "Dorado",
        "30cm",
        "3mm",
        "5.50",
        "2.00",
        "Feliz Cumpleaños",
        "50",
        "TRUE"

    ])

    format_sheet(ws)

    wb.save(path)


# ==========================================
# IMPORT CLIENTS
# ==========================================

def import_clients(db, file_path):

    wb = load_workbook(file_path)

    ws = wb.active

    imported = 0

    for row in ws.iter_rows(
        min_row=3,
        values_only=True
    ):

        if not row[0]:
            continue

        client = Client(

            name=row[0],
            company=row[1],
            ruc_ci=row[2],
            phone=row[3],
            email=row[4],
            address=row[5],
            client_type=row[6],
            observations=row[7]

        )

        db.add(client)

        imported += 1

    db.commit()

    return imported


# ==========================================
# IMPORT PRODUCTS
# ==========================================

def import_products(db, file_path):

    wb = load_workbook(file_path)

    ws = wb.active

    imported = 0

    for row in ws.iter_rows(
        min_row=3,
        values_only=True
    ):

        if not row[1]:
            continue

        product = Product(

            code=row[0],
            name=row[1],
            description=row[2],
            category=row[3],
            material=row[4],
            color=row[5],
            size=row[6],
            thickness=row[7],
            price=float(row[8] or 0),
            cost=float(row[9] or 0),
            theme=row[10],
            stock=int(row[11] or 0),
            custom=str(row[12]).lower() == "true"

        )

        db.add(product)

        imported += 1

    db.commit()

    return imported